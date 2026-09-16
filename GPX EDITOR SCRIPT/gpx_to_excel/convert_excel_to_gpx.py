#!/usr/bin/env python3
"""
convert_excel_to_gpx.py

Convert the Excel (.xlsx) file produced by convert_gpx_to_excel.py back into a
.gpx track file, in the same format/structure as a typical GPX 1.1 track
(<gpx><trk><trkseg><trkpt lat=".." lon=".."><ele>..</ele><time>..</time>
</trkpt>...).

Usage:
    python convert_excel_to_gpx.py

A window will pop up asking you to pick the .xlsx file to convert, then
where to save the resulting .gpx file. (You can still pass paths on the
command line instead if you prefer: `python convert_excel_to_gpx.py input.xlsx
[output.gpx]` skips the dialogs.)

Expected columns in the worksheet (matched by header text, case-insensitive,
any order): "Latitude", "Longitude", optionally "Elevation (m)" and
"Time (UTC)". Extra columns (Interval, Elapsed Time, Cum. Distance) are
ignored — they're derived data, not needed to rebuild the track.

Dependencies:
    pip install openpyxl
"""

import os
import sys
from datetime import datetime, timezone
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


def find_column_indices(header_row):
    """Map the columns we care about to their 1-based index, by header text."""
    wanted = {
        "latitude": None,
        "longitude": None,
        "elevation": None,
        "time": None,
    }
    for idx, cell_value in enumerate(header_row, start=1):
        if cell_value is None:
            continue
        text = str(cell_value).strip().lower()
        if text.startswith("lat"):
            wanted["latitude"] = idx
        elif text.startswith("lon"):
            wanted["longitude"] = idx
        elif text.startswith("elevation"):
            wanted["elevation"] = idx
        elif text.startswith("time"):
            wanted["time"] = idx

    missing = [k for k in ("latitude", "longitude") if wanted[k] is None]
    if missing:
        raise ValueError(
            f"Could not find required column(s) in the sheet: {', '.join(missing)}. "
            "Expected headers like 'Latitude' and 'Longitude'."
        )
    return wanted


def read_points_from_excel(path):
    wb = load_workbook(path, data_only=True)
    # Prefer a sheet literally named "GPX Track" if present (matches convert_gpx_to_excel.py output);
    # otherwise fall back to the first/active sheet.
    if "GPX Track" in wb.sheetnames:
        ws = wb["GPX Track"]
    else:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The worksheet is empty.")

    header_row = rows[0]
    cols = find_column_indices(header_row)

    points = []
    for row in rows[1:]:
        if row is None or all(v is None for v in row):
            continue

        lat = row[cols["latitude"] - 1]
        lon = row[cols["longitude"] - 1]
        if lat is None or lon is None:
            continue  # skip incomplete rows

        ele = row[cols["elevation"] - 1] if cols["elevation"] else None

        time_val = None
        if cols["time"]:
            raw_time = row[cols["time"] - 1]
            if isinstance(raw_time, datetime):
                time_val = raw_time
            elif isinstance(raw_time, str) and raw_time.strip():
                try:
                    time_val = datetime.fromisoformat(raw_time.strip())
                except ValueError:
                    time_val = None

        points.append({
            "lat": float(lat),
            "lon": float(lon),
            "ele": float(ele) if ele is not None else None,
            "time": time_val,
        })

    return points


def build_gpx(points, track_name="Track"):
    ns = "http://www.topografix.com/GPX/1/1"
    ET.register_namespace("", ns)

    gpx = ET.Element("gpx", {
        "version": "1.1",
        "creator": "convert_excel_to_gpx.py",
        "xmlns": ns,
    })
    trk = ET.SubElement(gpx, "trk")
    name_el = ET.SubElement(trk, "name")
    name_el.text = track_name
    trkseg = ET.SubElement(trk, "trkseg")

    for p in points:
        trkpt = ET.SubElement(trkseg, "trkpt", {
            "lat": repr(p["lat"]),
            "lon": repr(p["lon"]),
        })
        if p["ele"] is not None:
            ele_el = ET.SubElement(trkpt, "ele")
            ele_el.text = str(p["ele"])
        if p["time"] is not None:
            t = p["time"]
            if t.tzinfo is not None:
                t = t.astimezone(timezone.utc).replace(tzinfo=None)
            time_el = ET.SubElement(trkpt, "time")
            time_el.text = t.strftime("%Y-%m-%dT%H:%M:%SZ")

    tree = ET.ElementTree(gpx)
    ET.indent(tree, space="  ", level=0)  # Python 3.9+: pretty-print
    return tree


def pick_files_with_gui():
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    input_path = filedialog.askopenfilename(
        title="Select an Excel file to convert",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        parent=root,
    )
    if not input_path:
        root.destroy()
        return None, None

    default_name = os.path.splitext(os.path.basename(input_path))[0] + ".gpx"
    output_path = filedialog.asksaveasfilename(
        title="Save GPX file as",
        defaultextension=".gpx",
        initialfile=default_name,
        filetypes=[("GPX files", "*.gpx")],
        parent=root,
    )
    if not output_path:
        root.destroy()
        return None, None

    root.destroy()
    return input_path, output_path


def show_gui_message(title, text, error=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        if error:
            messagebox.showerror(title, text, parent=root)
        else:
            messagebox.showinfo(title, text, parent=root)
        root.destroy()
    except Exception:
        pass


def main():
    args = sys.argv[1:]
    if args:
        input_path = args[0]
        output_path = args[1] if len(args) > 1 else input_path.rsplit(".", 1)[0] + ".gpx"
    else:
        input_path, output_path = pick_files_with_gui()
        if not input_path:
            print("No file selected. Exiting.")
            return

    try:
        points = read_points_from_excel(input_path)
    except Exception as e:
        msg = f"Could not read Excel file:\n{e}"
        print(msg, file=sys.stderr)
        show_gui_message("Error", msg, error=True)
        sys.exit(1)

    if not points:
        msg = "No valid latitude/longitude rows found in the spreadsheet."
        print(msg, file=sys.stderr)
        show_gui_message("Error", msg, error=True)
        sys.exit(1)

    track_name = os.path.splitext(os.path.basename(output_path))[0]
    tree = build_gpx(points, track_name=track_name)
    tree.write(output_path, encoding="UTF-8", xml_declaration=True)

    msg = f"Wrote {len(points)} points to:\n{output_path}"
    print(msg)
    show_gui_message("Done", msg)


if __name__ == "__main__":
    main()