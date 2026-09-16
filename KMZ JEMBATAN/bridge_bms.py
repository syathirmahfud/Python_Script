#!/usr/bin/env python3
"""Bridge inventory generator with optional PKRMS Excel import.
Script version: V1.0 (Python script revision, not a BMS standard version).
Install: py -m pip install lxml Pillow reportlab pypdf openpyxl
Run: py bridge_bms_v1.0.py
Tkinter asks for the KMZ, data source (KMZ or PKRMS Excel), and output folder.
CLI: py bridge_bms_v1.0.py "test.kmz" --output-dir "results"
Excel: py bridge_bms_v1.0.py "test.kmz" --pkrms "INPUT_PKRMS_JEMBATAN.xlsx"
Use --pkrms without a path to read DEFAULT_PKRMS_PATH below.
Options: --kmz-only --overwrite --title "Pengabuan / 2026"

KMZ mode uses the existing verified popup table. Excel mode matches Bridge_Number
to the table ID, or to a raw bridge Point's name, ignoring case/invisible formatting.
Excel columns replace matched survey fields; absent columns retain existing values.
Unmatched bridges retain their existing attributes and are reported in the terminal.
Excel mode keeps map Point coordinates and photos from the KMZ. All Excel columns
are retained in ExtendedData on matched bridges; the workbook is read-only.
Blank cells, absent optional columns, -, en/em dashes, N/A, NA, None,
null and NaN are treated as unavailable and displayed as -; never as zero.
STA is formatted only when available. Missing table coordinates retain
the complete existing KMZ Point location (no mixing partial coordinate pairs).
Malformed nonempty numbers still report the bridge ID and field.
A bridge ID and a usable Point location remain necessary.
Legacy notes stay archived, with no length comparisons.
All embedded photos and timestamps are preserved. openpyxl is needed only for Excel mode.
Output photo filenames use "<bridge number> - (n)<original extension>",
starting at 1 for each bridge. Popup and archived photo links are updated.
Photo bytes, source order and the original input KMZ remain unchanged.
PKRMS code mappings and condition offsets follow the supplied enrichment script.
Missing condition inputs make the affected label unavailable, not "Baik".
Available inputs use its worst-condition/clamping rule; no NK BMS score is inferred.
Inputs are never overwritten; existing outputs require confirmation.
Failed saves restore earlier outputs; any failed recovery reports its backup folder.
Invisible formatting characters are ignored in IDs, field labels and numbers.
Source strings remain preserved in the KMZ; PDF IDs validate across text wrapping.
Supports KML 2.2, doc.kml or one KML, and two-column bridge popup tables.
Photos must be embedded; missing or corrupt images are reported.
Photos pass both format verification and full pixel decoding before output.
Bridge numbers starting with G (case-insensitive) are excluded from the
PDF register, cards, photos and totals, but retained in the styled KMZ.
Condition colours: Baik green, Sedang yellow, Rusak Ringan orange,
Rusak Berat red; unavailable or unrecognised labels grey. Labels remain intact.
Header ignores placeholder district/year values and deduplicates whitespace/case.
PDF: all condition highlighting appears only in the opening summary/register.
Bridge cards and component tables use neutral styling for every condition.
PDF includes only the first six photos per bridge, in original source order.
Its photo total counts displayed photos; the KMZ keeps all original photos.
PDF photo copies use EXIF orientation and a 300-DPI target without upscaling.
Resampled photos use JPEG quality 82; transparency is composited onto white.
PDF and KMZ popup headers include PU and Tanjabbar logos from the paths below.
Override with --pu-logo and --tanjabbar-logo; --no-logos explicitly disables them.
Missing or unreadable logo files produce an actionable error before output writes.
KMZ detail panels are neutral; condition colours remain on map icons.
"""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import ExitStack
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import html
import io
from math import ceil
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import sys
import tempfile
import logging
import time
import unicodedata
from urllib.parse import quote, unquote, urlsplit
import zipfile

try:
    from lxml import etree as ET
    from PIL import Image, ImageDraw, ImageOps
except ImportError as exc:
    raise SystemExit('Missing dependency. Run: py -m pip install lxml Pillow reportlab pypdf') from exc

KML = 'http://www.opengis.net/kml/2.2'
PDF_PHOTO_DPI = 300
PDF_PHOTO_JPEG_QUALITY = 82
LOG = logging.getLogger('bridge_bms')
LOG.addHandler(logging.NullHandler())
CURRENT_CONTEXT = 'Startup'

def placemark_details(pm):
    name = pm.findtext('k:name', default='Unnamed', namespaces=NS)
    folders = []
    for parent in reversed(list(pm.iterancestors())):
        if ET.QName(parent).localname in ('Folder', 'Document'):
            label = parent.findtext('k:name', default='', namespaces=NS)
            if label:
                folders.append(label)
    return f'placemark={name!r}; folder={"/".join(folders) or "(root)"}'

def context(stage, pm=None, ident=None, detail=''):
    global CURRENT_CONTEXT
    CURRENT_CONTEXT = stage
    if pm is not None:
        CURRENT_CONTEXT += ' | ' + placemark_details(pm)
    if ident is not None:
        CURRENT_CONTEXT += f' | bridge ID={ident!r}'
    if detail:
        CURRENT_CONTEXT += ' | ' + detail

def failure_details(exc):
    return f'{type(exc).__name__}: {exc}\nContext: {CURRENT_CONTEXT}'

def configure_logging():
    # stderr works in PowerShell/cmd and with Tee-Object. pythonw has no console.
    if not any(isinstance(h, logging.StreamHandler) for h in LOG.handlers):
        stream = sys.stderr
        if stream is not None:
            if hasattr(stream, 'reconfigure'):
                stream.reconfigure(errors='backslashreplace')
            handler = logging.StreamHandler(stream)
            handler.setFormatter(logging.Formatter('[%(asctime)s] %(levelname)s %(message)s', '%H:%M:%S'))
            LOG.addHandler(handler)
    LOG.setLevel(logging.INFO)
NS = {'k': KML}
NAVY, TEAL, AMBER, GREY, LIGHT = '#122D42', '#007E78', '#B16B00', '#526676', '#EDF3F6'
PALETTE = {'Baik': '#00FF0D', 'Sedang': '#F9EE25',
           'Rusak Ringan': '#FF8000', 'Rusak Berat': '#FF0000'}
PU_LOGO = Path(r'D:\\DEV\\assets\\logo_pupr.png')
TANJABBAR_LOGO = Path(r'D:\\DEV\\assets\\logo_TanjabBarat_2048.png')
DEFAULT_PKRMS_PATH = Path(r'D:\15.09 TANJUNG JABUNG BARAT\SURVEY KONDISI JEMBATAN\DOKUMEN\INPUT_PKRMS_JEMBATAN.xlsx')

# Mappings copied from the user's enrich_bridge_kmz_from_pkrms_excel.py.
CODE_BRIDGE_TYPE = {
    1: 'Gorong-gorong Bulat', 2: 'Gorong-gorong Oval', 3: 'Gorong-gorong Kotak',
    4: 'Oval', 5: 'Gelagar', 6: 'Balok Oval', 7: 'Komposit', 8: 'Plat',
    9: 'Rangka Baja', 10: 'Jembatan Gantung', 11: 'Jembatan Kabel Pancang',
    12: 'Jembatan Sementara', 13: 'Lintasan Kereta', 14: 'Lintasan Sungai',
    15: 'Lintasan Feri', 16: 'Lainnya', 17: 'Struktur Diperlukan',
}
CODE_BRIDGE_TYPE_SHORTCODE = dict(enumerate(
    ('YTI', 'API', 'BTI', 'EMI', 'GBI', 'LLI', 'MLI', 'PTI', 'RBA', 'TBI',
     'CBI', 'SBW', 'KLI', 'WTI', 'FLI', 'ULL', 'XXX'), 1))
CODE_PONDASI_TYPE = dict(enumerate(
    ('Cakar Ayam', 'Langsung', 'Tiang Pancang', 'Tiang Bor', 'Tiang Ulir', 'Sumuran', 'Lainnya'), 1))
CODE_BANGBWAH_TYPE = {1: 'Cap (Kepala Tiang)', 2: 'Dinding Penuh', 3: 'Kepala Jembatan Khusus'}
CODE_MATERIAL = dict(enumerate((
    'Kayu', 'Pasangan Bata', 'Pasangan Batu', 'Bronjong Dan Sejenisnya',
    'Pasangan Batu Kosong', 'Beton Tak Bertulang', 'Beton Bertulang',
    'Beton Pratekan', 'Baja', 'Plat Baja Bergelombang', 'Komposit Baja-Beton',
    'Aluminium', 'Neoprene/Karet', 'Teflon', 'PVC', 'Geotextile',
    'Tanah Biasa/Lempung atau Timbunan', 'Aspal', 'Kerikil', 'Macadam', 'Bahan Asli', 'Lain-Lain'), 1))
CODE_LANTAI_TYPE = {**CODE_MATERIAL, 6: 'Beton Bertulang'}
PKRMS_DIRECT = {
    'STA(m)': 'Chainage', 'Panjang(m)': 'Bridge_Length', 'Lebar(m)': 'Road_Width',
    'Jumlah Bentang': 'Number_Spans', 'Tahun Konstruksi': 'Year_Construction', 'Tahun Survey': 'Year',
}
PKRMS_CODE_FIELDS = {
    'Bangunan Atas - Kode': ('Bridge_Type', CODE_BRIDGE_TYPE_SHORTCODE),
    'Bangunan Atas - Tipe': ('Bridge_Type', CODE_BRIDGE_TYPE),
    'Bangunan Bawah - Tipe': ('BangbwahType', CODE_BANGBWAH_TYPE),
    'Bangunan Bawah - Bahan': ('BahanBangbwah', CODE_MATERIAL),
    'Fondasi - Tipe': ('PondasiType', CODE_PONDASI_TYPE),
    'Fondasi - Bahan': ('BahanPond', CODE_MATERIAL),
    'Permukaan Jembatan - Tipe': ('LantaiType', CODE_LANTAI_TYPE),
}
PKRMS_CONDITIONS = {
    'Bangunan Atas - Kondisi': (('Cond_DeckJoints', 1), ('Cond_Beam', 0)),
    'Bangunan Bawah - Kondisi': (('Cond_Piers', 0), ('Cond_Bearings', 0),
                               ('Cond_WingWalls', 1), ('Cond_Abutment', 0)),
    'Fondasi - Kondisi': (('Cond_Foundations', 0),),
    'Permukaan Jembatan - Kondisi': (('Cond_RoadSurface', 2), ('Cond_Deck', 1)),
    'Kondisi Jembatan (Keseluruhan)': (
        ('Cond_Piers', 0), ('Cond_Bearings', 0), ('Cond_Foundations', 0), ('Cond_Scouring', 1),
        ('Cond_RoadSurface', 2), ('Cond_Deck', 1), ('Cond_DeckJoints', 1), ('Cond_Beam', 0),
        ('Cond_WingWalls', 1), ('Cond_Abutment', 0)),
}
BRIDGE_NAME_PATTERN = re.compile(r'^([GJ]-\d+(?:\.\d+)*-\d+)(?:\s+(.*))?$', re.IGNORECASE)

def load_logos(paths):
    result = []
    for path in paths:
        try:
            source_data = Path(path).read_bytes()
            # Some Windows logo files have a .png extension but contain JPEG,
            # WEBP, or another valid raster format. Normalize every readable
            # logo to PNG for reliable ReportLab and KMZ embedding.
            with Image.open(io.BytesIO(source_data)) as im:
                im.load()
                w, h = im.size
                normalized = im.convert('RGBA')
                buf = io.BytesIO()
                normalized.save(buf, format='PNG', optimize=True)
                data = buf.getvalue()
        except (OSError, ValueError) as exc:
            raise ValueError(f'Cannot read logo {path}: {exc}. Correct the logo path or use --no-logos.') from exc
        result.append((data, w, h))
    return result
REQUIRED = ('No. Jembatan', 'Nama Jembatan', 'Longitude', 'Latitude',
            'STA(m)', 'Panjang(m)', 'Lebar(m)', 'Jumlah Bentang',
            'Kondisi Jembatan (Keseluruhan)')
GENERATED_ROWS = {'Catatan asli (Description)', 'Status penilaian', 'Sumber nilai', 'STA tampilan'}
POLICY = ('Data acuan: tabel KMZ atau Excel PKRMS; sumber dicantumkan pada setiap kartu. '
          'Nilai tidak tersedia ditampilkan sebagai tanda hubung, bukan nol. '
          'Catatan lama diarsipkan tanpa perbandingan panjang. NK BMS tidak dihitung.')
MISSING = {'', '-', '–', '—', 'n/a', 'na', 'none', 'null', 'nan'}

def clean_text(value):
    """Remove nonprinting format characters for interpretation and PDF display."""
    return ''.join(c for c in str(value if value is not None else '')
                   if unicodedata.category(c) != 'Cf')

def is_missing(value):
    return clean_text(value).strip().casefold() in MISSING

def bridge_key(value):
    return clean_text(value).strip().upper()

def pdf_has_bridge_id(pages, ident):
    # PDF extraction may split an ID across lines. Require the complete ID,
    # so J-...-1 cannot be satisfied by J-...-10 or another longer identifier.
    key = re.sub(r'\s+', '', bridge_key(ident))
    if not key:
        return False
    pattern = re.compile(r'(?<![\w.-])' + r'\s*'.join(map(re.escape, key))
                         + r'(?![\w.-])')
    return any(pattern.search(bridge_key(page)) for page in pages)

def display(value):
    return '-' if is_missing(value) else str(value).strip()

def condition_name(value):
    normalized = re.sub(r'[\s_\-]+', ' ', clean_text(value).strip()).casefold()
    return next((k for k in PALETTE if k.casefold() == normalized), display(value))

def condition_color(value):
    return PALETTE.get(condition_name(value), GREY)

def condition_ink(value):
    return NAVY if condition_name(value) in ('Baik', 'Sedang', 'Rusak Ringan') else '#FFFFFF'

def dossier_bridges(bridges):
    return [b for b in bridges if not bridge_key(b.fields['No. Jembatan']).startswith('G')]

def make_title(bridges):
    def unique(field):
        seen, result = set(), []
        for b in bridges:
            v = re.sub(r'\s+', ' ', clean_text(b.fields.get(field))).strip()
            if not is_missing(v) and v.casefold() not in seen:
                seen.add(v.casefold()); result.append(v)
        return result
    regions, years = unique('Kecamatan'), unique('Tahun Survey')
    return ' / '.join((regions or ['Inventaris jembatan']) + years)

def canonical_field(value):
    # Ignore harmless HTML whitespace, capitalisation and trailing colons.
    key = re.sub(r'\s+', ' ', clean_text(value)).strip().rstrip(':').strip()
    lookup = {k.casefold(): k for k in (*REQUIRED, 'Kecamatan', 'Tahun Survey')}
    lookup['tahun survei'] = 'Tahun Survey'
    return lookup.get(key.casefold(), key)

def field_number(value, ident, field):
    try:
        return number(value)
    except ValueError as exc:
        raise ValueError(f'{ident} / {field}: {exc}') from exc

def q(name):
    return '{' + KML + '}' + name

def escape(value):
    return html.escape(str(value), quote=True)

def number(value):
    if is_missing(value):
        return None
    s = clean_text(value).strip().replace(',', '.')
    try:
        n = Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f'Invalid number: {value!r}') from exc
    if not n.is_finite():
        raise ValueError(f'Non-finite number: {value!r}')
    return n

def sta_number(value):
    if is_missing(value):
        return None
    s = clean_text(value).strip().replace(',', '.')
    if '+' in s:
        if not re.fullmatch(r'\d+\+\d{1,3}(?:\.\d+)?', s):
            raise ValueError(f'Invalid station: {value!r}')
        km, m = s.split('+')
        return number(km) * 1000 + number(m)
    return number(s)

def station(value):
    n = sta_number(value)
    if n is None:
        return '-'
    if n < 0:
        raise ValueError('STA must not be negative.')
    km = int(n // 1000)
    remainder = format(n % 1000, 'f')
    whole, dot, fraction = remainder.partition('.')
    fraction = fraction.rstrip('0')
    return f'{km}+{int(whole):03d}' + ('.' + fraction if fraction else '')

INVENTORY_FIELDS = (
    'No. Jembatan', 'Nama Jembatan', 'Longitude', 'Latitude', 'Kecamatan',
    'STA(m)', 'Panjang(m)', 'Lebar(m)', 'Jumlah Bentang',
    'Bangunan Atas - Kode', 'Bangunan Atas - Tipe', 'Bangunan Atas - Kondisi',
    'Bangunan Bawah - Tipe', 'Bangunan Bawah - Bahan', 'Bangunan Bawah - Kondisi',
    'Fondasi - Tipe', 'Fondasi - Bahan', 'Fondasi - Kondisi',
    'Permukaan Jembatan - Tipe', 'Permukaan Jembatan - Kondisi',
    'Tahun Konstruksi', 'Kondisi Jembatan (Keseluruhan)', 'Tahun Survey',
)
PKRMS_RAW_PREFIX = 'pkrms_raw:'

@dataclass
class PKRMSWorkbook:
    path: Path
    sheet: str
    records: dict  # normalized Bridge_Number -> (Excel row number, all column values)

def pkrms_header(value):
    return re.sub(r'[\s_]+', '', clean_text(value)).casefold()

def index_pkrms_rows(rows, path, sheet):
    """Validate identifiers/headers before making any change to bridge data."""
    rows = iter(rows)
    raw_headers = next(rows, ())
    known = {'Bridge_Number', 'Bridge_Name', *PKRMS_DIRECT.values(),
             *(col for col, _ in PKRMS_CODE_FIELDS.values()),
             *(col for group in PKRMS_CONDITIONS.values() for col, _ in group)}
    canonical = {pkrms_header(col): col for col in known}
    headers, seen = [], set()
    for value in raw_headers:
        label = clean_text(value).strip()
        key = pkrms_header(label)
        if key and key in seen:
            raise ValueError(f'{path.name} / {sheet}: duplicate Excel column {label!r}.')
        seen.add(key)
        headers.append(canonical.get(key, label))
    if 'Bridge_Number' not in headers:
        raise ValueError(f'{path.name} / {sheet}: first row must contain Bridge_Number.')
    records = {}
    for row_number, row in enumerate(rows, 2):
        context('Reading PKRMS Excel row', detail=f'{path.name} / {sheet} / row={row_number}')
        if all(is_missing(value) for value in row):
            continue
        record = {key: value for key, value in zip(headers, row) if key}
        ident = bridge_key(record.get('Bridge_Number'))
        if is_missing(ident):
            raise ValueError(f'{path.name} / {sheet} / row {row_number}: missing Bridge_Number.')
        if ident in records:
            raise ValueError(f'{path.name} / {sheet}: duplicate Bridge_Number {ident!r} '
                             f'at rows {records[ident][0]} and {row_number}.')
        records[ident] = (row_number, record)
    if not records:
        raise ValueError(f'{path.name} / {sheet}: no PKRMS bridge rows found.')
    LOG.info('EXCEL: %d bridge records | %s / %s', len(records), path.name, sheet)
    return PKRMSWorkbook(path, sheet, records)

def load_pkrms_data(path):
    """Read the first worksheet, including cached formula values; never save it."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError('Excel import needs openpyxl. Run: py -m pip install openpyxl') from exc
    path = Path(path).expanduser().resolve()
    context('Opening PKRMS Excel', detail=str(path))
    if not path.is_file():
        raise ValueError(f'PKRMS Excel file not found: {path}')
    LOG.info('EXCEL: opening %s', path)
    with ExitStack() as stack:
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        stack.callback(cached.close)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        stack.callback(formulas.close)
        sheet, original = cached.worksheets[0], formulas.worksheets[0]
        def rows():
            for cells, source_cells in zip(sheet.iter_rows(), original.iter_rows()):
                for cell, source in zip(cells, source_cells):
                    if cell.data_type == 'e':
                        raise ValueError(f'{path.name} / {sheet.title}!{cell.coordinate}: Excel error {cell.value}.')
                    if source.data_type == 'f' and cell.value is None:
                        raise ValueError(f'{path.name} / {sheet.title}!{source.coordinate}: '
                                         'formula has no cached value. Recalculate and save in Excel first.')
                yield tuple(cell.value for cell in cells)
        return index_pkrms_rows(rows(), path, sheet.title)

def pkrms_integer(value, ident, column):
    n = field_number(value, ident, 'PKRMS.' + column)
    if n is not None and n != n.to_integral_value():
        raise ValueError(f'{ident} / PKRMS.{column}: expected an integer code, got {value!r}.')
    return None if n is None else int(n)

def pkrms_condition(record, inputs, ident):
    """Use the supplied offsets/max rule only when all required cells are present."""
    adjusted, missing = [], []
    for column, offset in inputs:
        value = pkrms_integer(record.get(column), ident, column)
        if value is None:
            missing.append(column)
        elif value - offset < 6:
            adjusted.append(value - offset)
    if missing or not adjusted:
        return '-', missing
    result = max(1, max(adjusted))
    result = 4 if result == 5 else result
    return {1: 'Baik', 2: 'Sedang', 3: 'Rusak Ringan', 4: 'Rusak Berat'}[result], []

def apply_pkrms_fields(fields, record, ident):
    """Replace imported attributes; absent columns leave existing attributes intact."""
    values = dict(fields)
    if is_missing(values.get('Nama Jembatan')) and not is_missing(record.get('Bridge_Name')):
        values['Nama Jembatan'] = display(record['Bridge_Name'])
    for label, column in PKRMS_DIRECT.items():
        if column not in record:
            continue
        value = record[column]
        if column == 'Chainage':
            try:
                value = station(value)
            except ValueError as exc:
                raise ValueError(f'{ident} / PKRMS.Chainage: {exc}') from exc
        elif column in ('Year', 'Year_Construction'):
            value = pkrms_integer(value, ident, column)
            if column == 'Year_Construction' and value == 0:
                value = None
        # Other dimensions are validated by load_input after the import.
        values[label] = display(value)
    for label, (column, codes) in PKRMS_CODE_FIELDS.items():
        if column not in record:
            continue
        code = pkrms_integer(record[column], ident, column)
        if code is None:
            values[label] = '-'
        elif code in codes:
            values[label] = codes[code]
        else:
            values[label] = f'Kode {code} (belum dipetakan)'
            LOG.warning('EXCEL CODE: %s | %s=%s is not in the supplied lookup table', ident, column, code)
    missing, unavailable = set(), False
    for label, inputs in PKRMS_CONDITIONS.items():
        # A workbook without any of this group's columns does not replace that group.
        if not any(column in record for column, _ in inputs):
            continue
        values[label], absent = pkrms_condition(record, inputs, ident)
        missing.update(absent)
        unavailable |= values[label] == '-'
    values.pop('Catatan kondisi PKRMS', None)
    if missing:
        values['Catatan kondisi PKRMS'] = 'Isian kondisi tidak lengkap; label terkait tidak dihitung.'
        LOG.warning('EXCEL CONDITION: %s | missing inputs: %s; affected labels remain unavailable',
                    ident, ', '.join(sorted(missing)))
    elif unavailable:
        values['Catatan kondisi PKRMS'] = 'Sebagian data kondisi tidak tersedia/tidak berlaku.'
    values['Status PKRMS'] = 'Cocok berdasarkan nomor jembatan'
    return values

def split_bridge_name(pm):
    text = clean_text(pm.findtext('k:name', default='', namespaces=NS)).strip()
    match = BRIDGE_NAME_PATTERN.fullmatch(text)
    return (bridge_key(match[1]), display(match[2])) if match else (None, '-')

def district_from_kmz(pm):
    # Nearest matching folder wins, so a multi-district KMZ keeps local context.
    for parent in pm.iterancestors():
        if parent.tag not in (q('Folder'), q('Document')):
            continue
        title = clean_text(parent.findtext('k:name', default='', namespaces=NS)).strip()
        match = re.search(r'\bKECAMATAN\s+(.+?)(?:\s*/\s*\d{4})?$', title, re.IGNORECASE)
        if match:
            return match[1].strip().title()
        match = re.fullmatch(r'([^/]+?)\s*/\s*\d{4}', title)
        if match and not re.search(r'\b(RUAS|JALAN|ROAD|SURVEY|INVENTARIS)\b', match[1], re.IGNORECASE):
            return display(match[1])
    return '-'

def bridge_photos(pm, description, kml_name, payload):
    """Keep popup order and add metadata-only photos; do not count header logos."""
    hrefs = description.xpath('//img[not(@data-bms-logo)]/@src') if description is not None else []
    photos = [(href, photo_member(href, kml_name, payload)) for href in hrefs]
    known = {member for _, member in photos}
    for node in pm.xpath('.//k:SimpleData[@name="pdfmaps_photos"] | '
                         './/k:Data[@name="pdfmaps_photos"]/k:value', namespaces=NS):
        h = ET.HTML(node.text or '<html/>')
        if h is None:
            continue
        metadata = [(href, photo_member(href, kml_name, payload))
                    for href in h.xpath('//img[not(@data-bms-logo)]/@src')]
        # All entries from a metadata-only source keep their order and repeats.
        photos.extend(pair for pair in metadata if pair[1] not in known)
        known.update(member for _, member in metadata)
    return photos

@dataclass
class Bridge:
    pm: object
    fields: dict
    photos: list  # (current href, ZIP member), in original source order
    authority: str = 'Hasil Survey Lapangan'
    pkrms_record: dict | None = None
    imported: bool = False

def photo_member(href, kml_name, names):
    parts = urlsplit(href)
    if parts.scheme or parts.netloc or parts.query or parts.fragment:
        raise ValueError(f'Only embedded photo paths are supported: {href}')
    decoded = unquote(href).replace('\\', '/')
    p = PurePosixPath(kml_name).parent / decoded
    if p.is_absolute() or '..' in p.parts:
        raise ValueError(f'Unsafe photo path: {href}')
    # Try literal paths first: a real filename can itself contain a percent sign.
    literal = str(PurePosixPath(kml_name).parent / href)
    for candidate in (literal, str(p)):
        if candidate in names:
            return candidate
    raise ValueError(f'Missing embedded photo: {href}')

def rename_bridge_photos(root, payload, kml_name, bridges):
    """Rename output photos and their links without re-encoding image bytes."""
    sources = {member for b in bridges for _, member in b.photos}
    if not sources:
        return dict(payload)
    occupied = {name.casefold() for name in payload if name not in sources}
    renamed, global_map, plans = {}, {}, []
    for b in bridges:
        ident = bridge_key(b.fields['No. Jembatan'])
        stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', ident).strip(' .')
        context('Naming bridge photos', b.pm, ident)
        if b.photos and (not stem or len(stem) > 180):
            raise ValueError(f'{ident!r}: bridge number cannot form a usable photo filename.')
        owner_map, photos = {}, []
        for index, (_, member) in enumerate(b.photos, 1):
            old = PurePosixPath(member)
            extension = old.suffix
            if not re.fullmatch(r'\.[A-Za-z0-9]{1,8}', extension):
                with Image.open(io.BytesIO(payload[member])) as im:
                    extension = next((ext for ext, fmt in Image.registered_extensions().items()
                                      if fmt == im.format), None)
                if extension is None:
                    raise ValueError(f'{ident} / photo {member!r}: cannot determine an image extension.')
            target = str(old.with_name(f'{stem} - ({index}){extension}'))
            if target.casefold() in occupied:
                raise ValueError(f'{ident} / photo {member!r}: output filename collision: {target}')
            occupied.add(target.casefold())
            renamed[target] = payload[member]
            # Shared images get a separate filename for each bridge/occurrence.
            # Other references to a shared source use its first matching copy.
            owner_map.setdefault(member, target)
            global_map.setdefault(member, target)
            href = quote(posixpath.relpath(target, str(PurePosixPath(kml_name).parent)), safe='/()-._')
            photos.append((href, target))
            LOG.info('PHOTO NAME: %s | %d/%d | %s -> %s', ident, index, len(b.photos), member, target)
        plans.append((b, photos, owner_map))

    def rewrite_tree(tree, document_name, scoped=False):
        changed = False
        owners = {b.pm: {**global_map, **mapping} for b, _, mapping in plans} if scoped else {}
        for element in tree.iter():
            if not isinstance(element.tag, str):
                continue
            mapping = global_map
            if scoped:
                owner = next((ancestor for ancestor in element.iterancestors()
                              if ancestor in owners), None)
                mapping = owners.get(element, owners.get(owner, global_map))
            def reference(value):
                decoded = html.unescape(value).strip()
                try:
                    parts = urlsplit(decoded)
                except ValueError:
                    return None  # Unrelated prose is not necessarily a valid URI.
                if parts.scheme or parts.netloc or parts.query or parts.fragment:
                    return None
                parent = str(PurePosixPath(document_name).parent)
                # Match literal percent signs first, just as photo_member does.
                for ref in (decoded, unquote(decoded)):
                    member = posixpath.normpath(posixpath.join(parent, ref.replace('\\', '/')))
                    if member in mapping:
                        return quote(posixpath.relpath(mapping[member], parent), safe='/()-._')
                return None
            def replace_links(value):
                direct = reference(value)
                if direct is not None:
                    return direct
                # Preserve HTML fragments (including pdfmaps_photos) verbatim
                # except for src/href values; accept quoted and unquoted URLs.
                def attribute(match):
                    url = next(v for v in match.groups()[1:] if v is not None)
                    new = reference(url)
                    return match.group(0) if new is None else match.group(1) + '"' + escape(new) + '"'
                return re.sub(r'''(\b(?:src|href)\s*=\s*)(?:"([^"]*)"|'([^']*)'|([^\s<>"']+))''',
                              attribute, value, flags=re.IGNORECASE)
            if element.text:
                new = replace_links(element.text)
                if new != element.text:
                    element.text = ET.CDATA(new) if '<' in new else new
                    changed = True
            for attr, value in list(element.attrib.items()):
                if ET.QName(attr).localname in ('href', 'src'):
                    new = reference(value)
                    if new is not None and new != value:
                        element.set(attr, new)
                        changed = True
        return changed

    result = {name: data for name, data in payload.items() if name not in sources}
    result.update(renamed)
    rewrite_tree(root, kml_name, scoped=True)
    for name, data in payload.items():
        if name != kml_name and name.lower().endswith('.kml'):
            context('Updating embedded KML photo links', detail=name)
            tree = ET.fromstring(data, ET.XMLParser(resolve_entities=False, no_network=True))
            if rewrite_tree(tree, name):
                result[name] = ET.tostring(tree, xml_declaration=True, encoding='UTF-8')
    for b, photos, _ in plans:
        b.photos = photos
    LOG.info('PHOTO NAME: %d photo entries named across %d bridges', len(renamed), len(bridges))
    return result

def load_input(path, pkrms=None):
    context('Opening input KMZ', detail=str(path))
    LOG.info('READ: opening KMZ %s', path)
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        names = [i.filename for i in infos]
        if len(names) != len(set(names)):
            raise ValueError('Duplicate ZIP member names; repair the input first.')
        if sum(i.file_size for i in infos) > 1024 ** 3:
            raise ValueError('Input exceeds the 1 GiB uncompressed safety limit.')
        candidates = [n for n in names if n.lower().endswith('.kml')]
        kml_name = 'doc.kml' if 'doc.kml' in names else (candidates[0] if len(candidates) == 1 else None)
        if not kml_name:
            raise ValueError('Expected doc.kml or exactly one KML file.')
        payload = {}
        for i, n in enumerate(names, 1):
            payload[n] = z.read(n)
            if i == 1 or i % 25 == 0 or i == len(names):
                LOG.info('READ: archive files %d/%d', i, len(names))
    parser = ET.XMLParser(resolve_entities=False, no_network=True)
    root = ET.fromstring(payload[kml_name], parser)
    if root.tag != q('kml'):
        raise ValueError('Expected KML 2.2 namespace.')
    bridges, ids, skipped = [], set(), 0
    matched, unmatched = set(), []
    placemarks = root.findall('.//k:Placemark', NS)
    for pm_index, pm in enumerate(placemarks, 1):
        context('Parsing placemark', pm, detail=f'{pm_index}/{len(placemarks)}')
        LOG.info('PARSE: placemark %d/%d | %s', pm_index, len(placemarks),
                 pm.findtext('k:name', default='Unnamed', namespaces=NS))
        h = ET.HTML(pm.findtext('k:description', default='', namespaces=NS) or '<html/>')
        fields = {}
        for tr in (h.xpath('//tr') if h is not None else []):
            cells = tr.xpath('./td | ./th')
            if len(cells) != 2 or any(cell.xpath('.//img | .//table') for cell in cells):
                continue
            key, value = [''.join(cell.itertext()).strip() for cell in cells]
            key = canonical_field(key)
            if not key or key in GENERATED_ROWS:
                continue
            if key in fields and fields[key] != value:
                raise ValueError(f'Conflicting duplicate table field: {key}')
            fields[key] = value
        has_table = 'No. Jembatan' in fields
        name_id, name_label = split_bridge_name(pm)
        if not has_table:
            if pkrms is None or name_id is None or pm.find('k:Point', NS) is None:
                skipped += 1
                continue
            fields = dict.fromkeys(INVENTORY_FIELDS, '-')
            fields.update({'No. Jembatan': name_id, 'Nama Jembatan': name_label})
            # Preserve a raw source description before the styled popup replaces it.
            original = pm.findtext('k:description', default='', namespaces=NS)
            if original:
                ex = pm.find('k:ExtendedData', NS)
                if ex is None:
                    ex = ET.SubElement(pm, q('ExtendedData'))
                if ex.find('k:Data[@name="bms_original_description"]', NS) is None:
                    data = ET.SubElement(ex, q('Data'), name='bms_original_description')
                    set_child(data, 'value', ET.CDATA(original))
        elif pkrms is not None and is_missing(fields['No. Jembatan']) and name_id:
            fields['No. Jembatan'] = name_id
        if is_missing(fields['No. Jembatan']):
            name = pm.findtext('k:name', default='Unnamed placemark', namespaces=NS)
            raise ValueError(f'{name}: No. Jembatan is missing.')
        fields = {k: display(v) for k, v in fields.items()}
        for key in REQUIRED:
            fields.setdefault(key, '-')
        ident = bridge_key(fields['No. Jembatan'])
        if pkrms is not None and name_id and name_id != ident:
            LOG.warning('EXCEL MATCH: popup ID %s differs from name ID %s; using the popup ID', ident, name_id)
        context('Validating bridge fields', pm, ident)
        if ident in ids:
            raise ValueError(f'Duplicate bridge ID: {ident}')
        ids.add(ident)
        coords = pm.findtext('k:Point/k:coordinates', default='', namespaces=NS).strip().split(',')
        if len(coords) not in (2, 3):
            raise ValueError(f'{ident} / Point: expected longitude,latitude[,altitude].')
        old_lon = field_number(coords[0], ident, 'Point longitude')
        old_lat = field_number(coords[1], ident, 'Point latitude')
        if old_lon is None or old_lat is None or not (-180 <= old_lon <= 180 and -90 <= old_lat <= 90):
            raise ValueError(f'{ident} / Point: existing map location is missing or out of range.')
        if len(coords) == 3 and field_number(coords[2], ident, 'Point altitude') is None:
            raise ValueError(f'{ident} / Point altitude: invalid geometry; omit altitude if unavailable.')
        authority = pm.findtext('k:ExtendedData/k:Data[@name="bms_data_authority"]/k:value',
                                default='Hasil Survey Lapangan', namespaces=NS)
        source_record = {data.get('name')[len(PKRMS_RAW_PREFIX):]: data.findtext('k:value', default='-', namespaces=NS)
                         for data in pm.findall('k:ExtendedData/k:Data', NS)
                         if (data.get('name') or '').startswith(PKRMS_RAW_PREFIX)} or None
        imported = False
        if pkrms is not None:
            if is_missing(fields.get('Nama Jembatan')):
                fields['Nama Jembatan'] = name_label
            if is_missing(fields.get('Kecamatan')):
                fields['Kecamatan'] = district_from_kmz(pm)
            entry = pkrms.records.get(ident)
            if entry is not None:
                row_number, source_record = entry
                context('Importing PKRMS bridge fields', pm, ident,
                        f'{pkrms.path.name} / {pkrms.sheet} / row={row_number}')
                fields = apply_pkrms_fields(fields, source_record, ident)
                authority = (f'PKRMS Excel: {pkrms.path.name}; {pkrms.sheet}, baris {row_number}. '
                             'Kolom tanpa impor, lokasi dan foto: KMZ.')
                imported = True
                matched.add(ident)
                LOG.info('EXCEL MATCH: %s | %s / row %d', ident, pkrms.sheet, row_number)
            else:
                unmatched.append(ident)
                fields['Status PKRMS'] = 'Tidak ada pasangan; atribut KMZ dipertahankan'
                LOG.warning('EXCEL UNMATCHED: %s | %s | existing KMZ values retained',
                            ident, placemark_details(pm))
            # Imported/raw bridges use KMZ geometry, never Excel coordinates.
            # Existing unmatched tables retain all their original field values.
            if imported or not has_table:
                fields['Longitude'], fields['Latitude'] = coords[0].strip(), coords[1].strip()
                fields['Sumber Koordinat'] = 'Point KMZ asli'
        lon = field_number(fields['Longitude'], ident, 'Longitude')
        lat = field_number(fields['Latitude'], ident, 'Latitude')
        if (lon is not None and not -180 <= lon <= 180) or (lat is not None and not -90 <= lat <= 90):
            raise ValueError(f'{ident} / Longitude, Latitude: coordinates outside valid ranges.')
        if lon is None or lat is None:
            fields['Sumber Koordinat'] = 'Point KMZ asli; pasangan koordinat tabel tidak lengkap'
        try:
            station(fields['STA(m)'])
        except ValueError as exc:
            raise ValueError(f'{ident} / STA(m): {exc}') from exc
        for key in ('Panjang(m)', 'Lebar(m)', 'Jumlah Bentang'):
            n = field_number(fields[key], ident, key)
            if n is not None and (n < 0 or (key == 'Jumlah Bentang' and n != int(n))):
                raise ValueError(f'{ident} / {key}: invalid value {fields[key]!r}')
        photos = []
        context('Reading bridge photo links', pm, ident)
        for href, member in bridge_photos(pm, h, kml_name, payload):
            context('Checking embedded photo', pm, ident, f'photo={href!r}')
            try:
                with Image.open(io.BytesIO(payload[member])) as image:
                    image.verify()
                # verify() does not decode JPEG pixels; reopen to catch truncation.
                with Image.open(io.BytesIO(payload[member])) as image:
                    image.load()
            except (OSError, ValueError, SyntaxError, Image.DecompressionBombError) as exc:
                raise ValueError(f'{ident} / photo {member!r}: invalid image: {exc}') from exc
            photos.append((href, member))
        bridges.append(Bridge(pm, fields, photos, authority, source_record, imported))
        LOG.info('PARSE OK: %s | %d photos checked', ident, len(photos))
    if not bridges:
        raise ValueError('No bridge popup tables found. Expected No. Jembatan; '
                         'for raw KMZ bridge Points, select Excel import or use --pkrms.')
    if pkrms is not None:
        unused = set(pkrms.records) - matched
        LOG.info('EXCEL SUMMARY: %d matched; %d KMZ bridges unmatched; %d Excel records unused',
                 len(matched), len(unmatched), len(unused))
        if unused:
            LOG.warning('EXCEL UNUSED: %s', ', '.join(sorted(unused)))
    return root, payload, kml_name, bridges, skipped

def set_child(parent, name, value):
    el = parent.find('k:' + name, NS)
    if el is None:
        el = ET.SubElement(parent, q(name))
    el.text = value
    return el

def order_feature(element):
    # KML Feature common fields precede container contents / geometry.
    rank = {'name': 0, 'visibility': 1, 'open': 2, 'Snippet': 5, 'snippet': 5,
            'description': 6, 'LookAt': 7, 'Camera': 7, 'TimeSpan': 8, 'TimeStamp': 8,
            'styleUrl': 9, 'Style': 10, 'StyleMap': 10, 'Region': 11,
            'Metadata': 12, 'ExtendedData': 12, 'Schema': 13}
    children = list(element)
    # Comments and processing instructions are not element tags. Keep their
    # slots while reordering only the KML elements.
    ordered = iter(sorted((child for child in children if isinstance(child.tag, str)),
                          key=lambda e: rank.get(ET.QName(e).localname, 14)))
    for child in children:
        element.remove(child)
    for child in children:
        element.append(next(ordered) if isinstance(child.tag, str) else child)

def popup(b, title, logo_refs=()):
    d = b.fields
    status = d['Kondisi Jembatan (Keseluruhan)']
    logo_html = ''.join(f'<img data-bms-logo="1" src="{escape(ref)}" width="{round(w*min(42/w,48/h))}" '
                        f'height="{round(h*min(42/w,48/h))}" style="margin:3px"/>' for ref,w,h in logo_refs)
    parts = [f'<div style="font-family:Arial,sans-serif;width:560px;color:{NAVY}">',
             f'<div style="background:{NAVY};color:white;padding:16px"><table width="100%"><tr><td><small style="color:#FFFFFF;">{escape(title)}</small>',
             f'<h2 style="color:#FFFFFF;">{escape(d["Nama Jembatan"])}</h2>'
             f'<span style="color:#FFFFFF;">{escape(d["No. Jembatan"])}</span></td>'
             + (f'<td width="110" align="right"><div style="background:white;padding:3px">{logo_html}</div></td>' if logo_refs else '')
             + '</tr></table></div>',
             f'<div style="background:{LIGHT};padding:12px"><b>STA {station(d["STA(m)"])} | '
             f'{escape(d["Panjang(m)"])} x {escape(d["Lebar(m)"])} m | {escape(d["Jumlah Bentang"])} bentang</b><br/>',
             f'<span>Kondisi: {escape(status)}</span> | NK BMS: tidak dihitung</div>',
             '<p>Data acuan: ' + escape(b.authority) + '</p>',
             '<h3>01 / Identitas &amp; komponen</h3><table style="width:100%;border-collapse:collapse;font-size:12px">']
    for k, v in d.items():
        parts.append(f'<tr><td style="padding:6px;border-bottom:1px solid #dce5ea;width:42%;color:{GREY}">{escape(k)}</td>'
                     f'<td style="padding:6px;border-bottom:1px solid #dce5ea">{escape(v)}</td></tr>')
    parts.append('</table><h3>02 / Dokumentasi sumber</h3><table>')
    for j in range(0, len(b.photos), 2):
        parts.append('<tr>')
        for href, member in b.photos[j:j + 2]:
            parts.append(f'<td style="vertical-align:top;width:270px"><a href="{escape(href)}">'
                         f'<img src="{escape(href)}" width="260"/></a><br/><small>'
                         f'{escape(PurePosixPath(member).name)}</small></td>')
        parts.append('</tr>')
    parts.append('</table>')
    if not b.photos:
        parts.append('<p>Tidak ada foto pada sumber.</p>')
    parts.append('<p style="font-size:11px;color:#526676">' + escape(POLICY) +
                 ' Tata letak BMS-inspired; bukan formulir resmi.</p></div>')
    return ''.join(parts)

def write_kmz(path, root, payload, kml_name, bridges, title, logos=()):
    LOG.info('KMZ: building %d bridge panels', len(bridges))
    doc = root.find('k:Document', NS)
    if doc is None:
        raise ValueError('Expected a top-level KML Document.')
    used_ids = set(root.xpath('//@id'))
    prefix = 'verified_bms'
    while any(str(i).startswith(prefix) for i in used_ids) or any(prefix in n for n in payload):
        prefix += '_v'
    icons, styles = {}, {}
    logo_refs = []
    for i,(data,w,h) in enumerate(logos):
        href = f'{prefix}_logos/{i}.png'
        icons[str(PurePosixPath(kml_name).parent / href)] = data
        logo_refs.append((href,w,h))
    for i, status in enumerate(dict.fromkeys(b.fields['Kondisi Jembatan (Keseluruhan)'] for b in bridges)):
        sid = f'{prefix}_{i}'
        href = f'{prefix}_icons/{i}.png'
        member = str(PurePosixPath(kml_name).parent / href)
        im = Image.new('RGBA', (64, 64))
        draw = ImageDraw.Draw(im)
        draw.ellipse((4, 4, 60, 60), fill=condition_color(status), outline='white', width=4)
        draw.line([(15, 39), (15, 24), (49, 24), (49, 39)], fill='white', width=4)
        draw.line([(15, 32), (49, 32)], fill='white', width=4)
        for x in (26, 38):
            draw.line([(x, 25), (x, 38)], fill='white', width=3)
        buf = io.BytesIO(); im.save(buf, format='PNG'); icons[member] = buf.getvalue()
        style = ET.SubElement(doc, q('Style'), id=sid)
        icon_style = ET.SubElement(style, q('IconStyle'))
        set_child(icon_style, 'scale', '1.1')
        icon = ET.SubElement(icon_style, q('Icon')); set_child(icon, 'href', href)
        label = ET.SubElement(style, q('LabelStyle')); set_child(label, 'scale', '0.8')
        balloon = ET.SubElement(style, q('BalloonStyle'))
        set_child(balloon, 'text', ET.CDATA('$[description]'))
        styles[status] = sid
    set_child(doc, 'name', title)
    set_child(doc, 'description', ET.CDATA(POLICY + ' Hijau: Baik; kuning: Sedang; jingga: Rusak Ringan; merah: Rusak Berat; abu-abu: kondisi tidak tersedia/tidak dikenali. Warna editorial.'))
    updated = 0
    for bridge_index, b in enumerate(bridges, 1):
        context('Building KMZ panel', b.pm, b.fields['No. Jembatan'])
        LOG.info('KMZ: bridge %d/%d | %s', bridge_index, len(bridges), b.fields['No. Jembatan'])
        pm, d = b.pm, b.fields
        coords = pm.find('k:Point/k:coordinates', NS)
        old = (coords.text or '').strip().split(',')
        if len(old) not in (2, 3) or any(any(c.isspace() for c in x.strip()) for x in old):
            raise ValueError(f'{d["No. Jembatan"]}: malformed Point coordinate tuple.')
        lon, lat = number(d['Longitude']), number(d['Latitude'])
        if lon is not None and lat is not None and not b.imported:
            new = [format(lon, 'f'), format(lat, 'f')]
            updated += int(number(old[0]) != lon or number(old[1]) != lat)
            if len(old) == 3:
                new.append(old[2].strip())
            coords.text = ','.join(new)
        # Incomplete table coordinates: leave original geometry text untouched.
        set_child(pm, 'name', d['No. Jembatan'] + ' ' + d['Nama Jembatan'])
        set_child(pm, 'description', ET.CDATA(popup(b, title, logo_refs)))
        set_child(pm, 'styleUrl', '#' + styles[d['Kondisi Jembatan (Keseluruhan)']])
        # Inline IconStyle/BalloonStyle would override the shared design.
        for inline in pm.findall('k:Style', NS):
            for child in list(inline):
                if child.tag in (q('IconStyle'), q('BalloonStyle')):
                    inline.remove(child)
        ex = pm.find('k:ExtendedData', NS)
        if ex is None:
            ex = ET.SubElement(pm, q('ExtendedData'))
        values = {**d, 'STA_formatted': station(d['STA(m)']),
                  'bms_data_authority': b.authority,
                  'bms_nk_status': 'Not calculated'}
        if b.pkrms_record is not None:
            values.update({PKRMS_RAW_PREFIX + key: display(value) for key, value in b.pkrms_record.items()})
        for data in list(ex.findall('k:Data', NS)):
            if (data.get('name') in values or data.get('name') == 'qa_length_conflict'
                    or (b.imported and (data.get('name') or '').startswith(PKRMS_RAW_PREFIX))):
                ex.remove(data)
        for key, val in values.items():
            data = ET.Element(q('Data'), name=key)
            set_child(data, 'value', val)
            ex.insert(len(ex.findall('k:Data', NS)), data)
        order_feature(pm)
    order_feature(doc)
    LOG.info('KMZ: compressing archive')
    context('Compressing KMZ archive', detail=str(path))
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as dest:
        for name, data in payload.items():
            dest.writestr(name, ET.tostring(root, xml_declaration=True, encoding='UTF-8') if name == kml_name else data)
        for name, data in icons.items():
            dest.writestr(name, data)
    LOG.info('KMZ: checking CRC and embedded-file preservation')
    context('Validating KMZ archive', detail=str(path))
    with zipfile.ZipFile(path) as check:
        if check.testzip():
            raise ValueError('Output KMZ failed CRC validation.')
        for name, data in payload.items():
            if name != kml_name and check.read(name) != data:
                raise ValueError(f'Embedded file changed unexpectedly: {name}')
    return updated

def write_pdf(path, bridges, payload, title, source, logos=()):
    bridges = dossier_bridges(bridges)
    LOG.info('PDF: preparing %d bridges after G-prefix filtering', len(bridges))
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                   TableStyle, Image as RLImage, PageBreak, KeepTogether)
    from pypdf import PdfReader
    from reportlab.lib.utils import ImageReader
    width = A4[0] - 72
    styles = {name: ParagraphStyle(name, fontName='Helvetica-Bold' if bold else 'Helvetica',
                                   fontSize=size, leading=size * 1.35, textColor=colors.HexColor(color),
                                   spaceAfter=6, splitLongWords=True, alignment=TA_LEFT)
              for name, size, bold, color in [('body', 9, False, NAVY), ('small', 7, False, GREY),
                                               ('heading', 17, True, NAVY), ('section', 11, True, NAVY),
                                               ('white', 10, True, '#FFFFFF')]}
    def p(value, style='body'):
        return Paragraph(escape(clean_text(value)), styles[style])
    def table(rows, widths, header=False, repeat_first=False):
        cells = [[p(v, 'white' if header and i == 0 else 'body') for v in row] for i, row in enumerate(rows)]
        condition_cells = []
        for i, row in enumerate(rows):
            if header and i == 0:
                continue
            col = (1 if len(row) == 2 and 'kondisi' in str(row[0]).casefold()
                   else 5 if len(row) == 6
                   else 4 if len(row) == 5 else None)
            if col is not None:
                value = row[col]
                if len(row) == 2:
                    continue
                cells[i][col] = Paragraph(escape(clean_text(value)), ParagraphStyle('condition',fontName='Helvetica',
                    fontSize=9,leading=12,textColor=colors.HexColor(condition_ink(value))))
                condition_cells.append(('BACKGROUND',(col,i),(col,i),colors.HexColor(condition_color(value))))
        t = Table(cells, colWidths=widths, repeatRows=1 if header or repeat_first else 0, hAlign='LEFT')
        commands = [('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
                    ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.HexColor(LIGHT), colors.white])]
        if header:
            commands.append(('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(NAVY)))
        commands.extend(condition_cells)
        t.setStyle(TableStyle(commands))
        return t
    def photograph(member, max_width, max_height):
        buf = io.BytesIO()
        with Image.open(io.BytesIO(payload[member])) as source:
            with ImageOps.exif_transpose(source) as im:
                iw, ih = im.size
                scale = min(max_width / iw, max_height / ih)
                draw_width, draw_height = iw * scale, ih * scale
                # ReportLab uses points (72 per inch), not screen pixels.
                bounds = (max(1, ceil(draw_width * PDF_PHOTO_DPI / 72)),
                          max(1, ceil(draw_height * PDF_PHOTO_DPI / 72)))
                im.thumbnail(bounds, Image.Resampling.LANCZOS)
                if (source.format == 'JPEG' and source.mode in ('RGB', 'L')
                        and source.getexif().get(274, 1) == 1 and im.size == (iw, ih)):
                    # Avoid another lossy encoding when a JPEG already fits.
                    buf = io.BytesIO(payload[member])
                else:
                    with im.convert('RGBA') as rgba, Image.new('RGB', im.size, 'white') as rgb:
                        rgb.paste(rgba, mask=rgba.getchannel('A'))
                        rgb.save(buf, format='JPEG', quality=PDF_PHOTO_JPEG_QUALITY)
        buf.seek(0)
        return RLImage(buf, width=draw_width, height=draw_height)
    def page_decoration(c, doc):
        context('Rendering PDF', detail=f'page={doc.page}; see preceding PDF PREPARE logs for bridge details')
        LOG.info('PDF RENDER: page %d', doc.page)
        c.saveState()
        c.setFillColor(colors.HexColor(NAVY)); c.rect(0, A4[1] - 84, A4[0], 84, fill=1, stroke=0)
        heading = Paragraph(escape(clean_text(title).upper()), ParagraphStyle('header', fontName='Helvetica-Bold',
                            fontSize=10, leading=12, textColor=colors.HexColor('#73CCC0')))
        text_width = width - (110 if logos else 0)
        _, hh = heading.wrap(text_width, 26)
        if hh > 26:
            raise ValueError('Title too long; shorten --title.')
        heading.drawOn(c, 36, A4[1] - 18 - hh)
        font_size = min(21, 21*text_width/c.stringWidth('Dossier inventaris jembatan','Helvetica-Bold',21))
        c.setFillColor(colors.white); c.setFont('Helvetica-Bold', font_size)
        c.drawString(36, A4[1] - 65, 'Dossier inventaris jembatan')
        if logos:
            x = A4[0] - 36 - 100
            c.setFillColor(colors.white); c.roundRect(x, A4[1]-68, 100, 52, 3, fill=1, stroke=0)
            for i,(data,w,h) in enumerate(logos):
                scale = min(42/w,44/h)
                c.drawImage(ImageReader(io.BytesIO(data)),x+4+i*48+(42-w*scale)/2,
                            A4[1]-64+(44-h*scale)/2,w*scale,h*scale,mask='auto')
        c.setStrokeColor(colors.HexColor('#D7E1E7')); c.line(36, 38, A4[0] - 36, 38)
        c.setFillColor(colors.HexColor(GREY)); c.setFont('Helvetica', 7)
        c.drawString(36, 24, 'Survey Kondisi Jembatan Kabupaten Tanjung Jabung Barat 2026')
        c.drawRightString(A4[0] - 36, 24, f'{doc.page:02d}')
        c.restoreState()
    counts = Counter(condition_name(b.fields['Kondisi Jembatan (Keseluruhan)']) for b in bridges)
    stat_style = ParagraphStyle('stat', fontName='Helvetica-Bold', fontSize=36, leading=42,
                                textColor=colors.HexColor(NAVY))
    stats = Table([[Paragraph(f'{len(bridges):02d}', stat_style), p('jembatan'),
                    Paragraph(f'{sum(min(6, len(b.photos)) for b in bridges):02d}', stat_style), p('foto ditampilkan')]],
                  colWidths=[75, width/2-75, 85, width/2-85], hAlign='LEFT')
    stats.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                              ('LEFTPADDING',(0,0),(-1,-1),0), ('BOTTOMPADDING',(0,0),(-1,-1),12)]))
    legend_cells = []
    for label, color in PALETTE.items():
        legend_cells.append(Paragraph(escape(f'{label}: {counts.get(label, 0)}'),
            ParagraphStyle('legend', fontName='Helvetica-Bold', fontSize=8, leading=11,
                           textColor=colors.HexColor(condition_ink(label)))))
    legend = Table([legend_cells], colWidths=[width/4]*4, hAlign='LEFT')
    legend.setStyle(TableStyle([('BACKGROUND',(i,0),(i,0),colors.HexColor(col))
                               for i,col in enumerate(PALETTE.values())] +
                              [('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    story = [stats, legend, Spacer(1, 12), p('Sumber: ' + source),
             p(' | '.join(f'{k}: {v}' for k,v in counts.items() if k not in PALETTE), 'small'),
             Spacer(1, 10), p('01 / Register jembatan', 'section')]
    rows = [['No.', 'No. Jembatan', 'Nama jembatan', 'STA', 'P (m)', 'Kondisi']]
    rows.extend([str(i), bridge_key(b.fields['No. Jembatan']), b.fields['Nama Jembatan'], station(b.fields['STA(m)']),
                 b.fields['Panjang(m)'], b.fields['Kondisi Jembatan (Keseluruhan)']]
                for i, b in enumerate(bridges, 1))
    register = table(rows, [32, 125, width - 352, 65, 45, 85], True)
    story += [register, Spacer(1, 15),
              p('02 / Dasar data', 'section'), p(POLICY),
              p('PDF untuk laporan dan pencetakan; KMZ untuk peta, atribut, dan foto. '
                'Koordinat tabel dipakai jika lengkap; jika tidak, lokasi Point KMZ dipertahankan. Ketinggian dan timestamp tetap.'),
              p('Format ini tidak menyatakan kapasitas struktur atau keamanan operasional. '
                'Kode kerusakan, skor NK dan rekomendasi penanganan tidak disimpulkan dari foto.')]
    if bridges:
        story.append(PageBreak())
    else:
        story.append(p('Tidak ada jembatan untuk dossier setelah filter nomor berawalan G.'))
    for i, b in enumerate(bridges, 1):
        d = b.fields
        context('Preparing PDF bridge card', b.pm, d['No. Jembatan'])
        LOG.info('PDF PREPARE: bridge %d/%d | %s | %s', i, len(bridges), d['No. Jembatan'], d['Nama Jembatan'])
        status = d['Kondisi Jembatan (Keseluruhan)']
        story += [p(f'{i:02d} / Kartu inventaris', 'section'), p(d['Nama Jembatan'], 'heading'),
                  p(d['No. Jembatan'] + ' | STA ' + station(d['STA(m)']))]
        band_text = Paragraph(escape('KONDISI: ' + clean_text(status).upper() + ' / NK BMS: TIDAK DIHITUNG'),
            ParagraphStyle('band',fontName='Helvetica-Bold',fontSize=10,leading=14,
                           textColor=colors.HexColor(NAVY)))
        band = Table([[band_text]], colWidths=[width])
        band.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(LIGHT)),
                                 ('LEFTPADDING', (0, 0), (-1, -1), 9), ('TOPPADDING', (0, 0), (-1, -1), 7)]))
        story += [band, Spacer(1, 8)]
        # All original table fields, not a fixed subset; long tables paginate.
        ordered_rows = [('No. Jembatan', d['No. Jembatan'])] + [(k, v) for k, v in d.items() if k != 'No. Jembatan']
        story += [table(ordered_rows, [width * .40, width * .60], repeat_first=True), Spacer(1, 9),
                  p('Sumber nilai: ' + b.authority, 'small')]
        
        if not b.photos:
            story += [p('Tidak ada foto pada sumber.')]
        shown_photos = b.photos[:6]
        for offset in range(0, len(shown_photos), 6):
            story += [PageBreak(), p(f'{i:02d} / Dokumentasi foto', 'section'), p(d['Nama Jembatan'], 'heading'),
                      p(d['No. Jembatan'] + f' | Foto 1-{len(shown_photos)} ditampilkan dari {len(b.photos)} foto sumber', 'small')]
            chunk = shown_photos[offset:offset + 6]
            for j in range(0, len(chunk), 2):
                cells = []
                for k, (_, member) in enumerate(chunk[j:j + 2]):
                    context('Preparing PDF photo', b.pm, d['No. Jembatan'], f'photo={member!r}')
                    LOG.info('PDF PHOTO: %s | %d/%d | %s', d['No. Jembatan'],
                             offset+j+k+1, len(shown_photos), member)
                    cells.append([photograph(member, (width - 24) / 2, 155), Spacer(1, 3),
                                  p(f'{offset+j+k+1:02d} / {PurePosixPath(member).name}', 'small')])
                if len(cells) == 1:
                    cells.append('')
                t = Table([cells], colWidths=[width / 2] * 2, hAlign='LEFT')
                t.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                      ('LEFTPADDING', (0, 0), (-1, -1), 0),
                                      ('BOTTOMPADDING', (0, 0), (-1, -1), 12)]))
                story.append(t)
            story.append(p('Foto dan urutan mengikuti sumber. Arah foto dan kode kerusakan tidak direkayasa.', 'small'))
        if i != len(bridges):
            story.append(PageBreak())
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=36, rightMargin=36,
                            topMargin=105, bottomMargin=52, title=title, author='Diolah dari ' + source)
    LOG.info('PDF: rendering document (final page count not known yet)')
    doc.build(story, onFirstPage=page_decoration, onLaterPages=page_decoration)
    reader = PdfReader(str(path))
    extracted = []
    for page_index, page in enumerate(reader.pages, 1):
        context('Extracting PDF text', detail=f'page={page_index}')
        LOG.info('PDF VALIDATE: extracting page %d/%d', page_index, len(reader.pages))
        extracted.append(page.extract_text() or '')
    failures = []
    for i, b in enumerate(bridges, 1):
        ident = b.fields['No. Jembatan']
        context('Validating PDF bridge ID', b.pm, ident)
        if not pdf_has_bridge_id(extracted, ident):
            failures.append(f'bridge ID={ident!r}; '+placemark_details(b.pm))
            LOG.error('PDF VALIDATE: complete ID not found after normalization | %r | %s', ident, b.fields['Nama Jembatan'])
        else:
            LOG.info('PDF VALIDATE: bridge %d/%d OK | %s', i, len(bridges), ident)
    if failures:
        raise ValueError('PDF validation failed: missing bridge ID(s) in extracted text: '
                         + '\n' + '\n'.join(failures)
                         + '. IDs were checked after removing invisible formatting and allowing line wrapping; see terminal diagnostics.')
    LOG.info('PDF: validation passed | %d pages', len(reader.pages))
    return len(reader.pages)

def job_paths(args):
    """Resolve destinations without reading the archive (safe on the GUI thread)."""
    source = Path(args.input).expanduser().resolve()
    context('Checking input and output paths', detail=str(source))
    if not source.is_file():
        raise ValueError(f'Input file not found: {source}')
    out = Path(args.output_dir or source.parent / (source.stem + '_bms_output')).expanduser().resolve()
    targets = [out / (source.stem + '_BMS.kmz')]
    if not args.kmz_only:
        targets.append(out / (source.stem + '_Dossier.pdf'))
    for target in targets:
        if target.resolve() == source:
            raise ValueError('Refusing to overwrite input.')
    return source, out, targets

def prepare_job(args):
    source, out, targets = job_paths(args)
    pkrms_path = getattr(args, 'pkrms', None)
    pkrms = load_pkrms_data(pkrms_path) if pkrms_path is not None else None
    root, payload, kml_name, bridges, skipped = load_input(source, pkrms)
    title = args.title or make_title(bridges)
    return source, root, payload, kml_name, bridges, skipped, title, out, targets

def publish_outputs(staged, targets, overwrite):
    """Publish the validated pair, restoring previous files on a normal save error."""
    existing = [target for target in targets if target.exists()]
    for target in existing:
        if not overwrite:
            raise ValueError(f'Output exists: {target}. Use another --output-dir or --overwrite.')
        if not target.is_file():
            raise ValueError(f'Output path is not a file: {target}')
    # Keep backups outside the auto-deleted staging directory. If recovery also
    # fails (e.g. a locked Windows file), the originals remain available here.
    backup_dir = (Path(tempfile.mkdtemp(prefix='bridge_bms_backup_', dir=targets[0].parent))
                  if existing else None)
    backups, published, recovery_errors = {}, [], []
    try:
        for target in existing:
            context('Backing up previous output', detail=str(target))
            backup = backup_dir / target.name
            shutil.copy2(target, backup)
            backups[target] = backup
        for temporary, target in zip(staged, targets):
            context('Saving validated output', detail=str(target))
            LOG.info('SAVE: %s', target)
            if target.exists() and not overwrite:
                raise ValueError(f'Output appeared during processing: {target}')
            temporary.replace(target)
            published.append(target)
    except Exception as exc:
        for target in reversed(published):
            try:
                if target in backups:
                    backups[target].replace(target)
                else:
                    target.unlink()
            except OSError as recovery:
                recovery_errors.append(f'{target}: {recovery}')
        if recovery_errors:
            raise OSError(f'Save failed: {exc}. Recovery also failed: '
                          + '; '.join(recovery_errors)
                          + (f'. Previous outputs retained in: {backup_dir}' if backup_dir else '')
                          + '. Close files in other applications before retrying.') from exc
        raise
    finally:
        # Failed recovery deliberately leaves the backups for manual recovery.
        if backup_dir and not recovery_errors:
            try:
                shutil.rmtree(backup_dir)
            except OSError as exc:
                LOG.warning('Could not remove backup folder %s: %s', backup_dir, exc)

def execute_job(args, job):
    started = time.monotonic()
    source, root, payload, kml_name, bridges, skipped, title, out, targets = job
    LOG.info('ASSETS: loading header logos%s', ' (disabled)' if args.no_logos else '')
    context('Loading header logos')
    logos = [] if args.no_logos else load_logos([args.pu_logo, args.tanjabbar_logo])
    for target in targets:
        if target.exists() and not args.overwrite:
            raise ValueError(f'Output exists: {target}. Use another --output-dir or --overwrite.')
    out.mkdir(parents=True, exist_ok=True)
    # Build and validate in staging before publishing final files.
    with tempfile.TemporaryDirectory(prefix='bridge_bms_', dir=out) as tmp:
        staged = [Path(tmp) / target.name for target in targets]
        payload = rename_bridge_photos(root, payload, kml_name, bridges)
        updated = write_kmz(staged[0], root, payload, kml_name, bridges, title, logos)
        pdf_bridges = dossier_bridges(bridges)
        pdf_title = args.title or make_title(pdf_bridges)
        pages = None if args.kmz_only else write_pdf(staged[1], pdf_bridges, payload, pdf_title,
                                                    source.name, logos)
        publish_outputs(staged, targets, args.overwrite)
    LOG.info('DONE: all outputs saved | %.1f seconds', time.monotonic() - started)
    import_summary = ''
    if getattr(args, 'pkrms', None) is not None:
        matched_count = sum(b.imported for b in bridges)
        import_summary = (f'PKRMS: {matched_count} cocok; {len(bridges)-matched_count} tanpa pasangan '
                          '(atribut KMZ dipertahankan). Lihat rincian di terminal.\n')
    return (f'OK: {len(bridges)} jembatan; {sum(len(b.photos) for b in bridges)} foto.\n'
            f'{updated} koordinat diperbarui; {skipped} placemark non-jembatan dipertahankan.\n'
            + import_summary
            + (f'PDF: {pages} halaman; {len(pdf_bridges)} jembatan; {len(bridges)-len(pdf_bridges)} nomor berawalan G dikecualikan.\n' if pages is not None else '')
            + '\n'.join(str(target) for target in targets))

def run_gui(args):
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    import queue
    import threading
    win = tk.Tk(); win.withdraw()
    exit_code = 0
    try:
        source = filedialog.askopenfilename(parent=win, title='1. Pilih file KMZ jembatan',
                                            filetypes=[('KMZ', '*.kmz')])
        if not source:
            return 0
        if getattr(args, 'pkrms', None) is None:
            use_excel = messagebox.askyesnocancel(
                'Sumber data jembatan',
                'Impor atribut jembatan dari Excel PKRMS?\n\n'
                'Ya: pilih Excel; nomor yang cocok memakai atribut Excel.\n'
                'Tidak: gunakan tabel yang sudah ada di KMZ.\n'
                'Batal: tutup tanpa membuat hasil.', parent=win)
            if use_excel is None:
                return 0
            if use_excel:
                excel = filedialog.askopenfilename(
                    parent=win, title='2. Pilih Excel PKRMS',
                    initialdir=str(DEFAULT_PKRMS_PATH.parent if DEFAULT_PKRMS_PATH.parent.is_dir()
                                   else Path(source).parent),
                    initialfile=DEFAULT_PKRMS_PATH.name,
                    filetypes=[('PKRMS Excel', '*.xlsx')])
                if not excel:
                    return 0
                args.pkrms = Path(excel)
        destination = filedialog.askdirectory(parent=win, title='Pilih folder penyimpanan hasil',
                                              initialdir=str(Path(source).parent), mustexist=False)
        if not destination:
            return 0
        args.input, args.output_dir = source, Path(destination)
        _, _, targets = job_paths(args)
        existing = [p for p in targets if p.exists()]
        if existing and not args.overwrite:
            if not messagebox.askyesno('Hasil sudah ada', 'Ganti file hasil berikut?\n' +
                                      '\n'.join(str(p) for p in existing), parent=win):
                return 0
            args.overwrite = True
        win.title('BMS - Membuat KMZ dan PDF')
        win.geometry('540x150'); win.resizable(False, False)
        ttk.Label(win, text='Membaca data KMZ/Excel dan menyusun hasil...', padding=16).pack()
        progress = ttk.Progressbar(win, mode='indeterminate', length=480)
        progress.pack(padx=20, pady=8); progress.start(12)
        win.deiconify()
        results = queue.Queue()
        def worker():
            try:
                # Archive reads, image decoding and PDF generation all run after
                # the progress window opens, without blocking Tk's event loop.
                results.put((True, execute_job(args, prepare_job(args))))
            except Exception as exc:
                LOG.exception('FAILED: %s', failure_details(exc))
                results.put((False, failure_details(exc)))
        def poll():
            nonlocal exit_code
            try:
                ok, result = results.get_nowait()
            except queue.Empty:
                win.after(100, poll)
                return
            progress.stop(); win.withdraw()
            if ok:
                messagebox.showinfo('Selesai', result, parent=win)
            else:
                exit_code = 1
                messagebox.showerror('Gagal membuat hasil', result, parent=win)
            win.quit()
        # Do not terminate halfway through writing if the progress window closes.
        win.protocol('WM_DELETE_WINDOW', lambda: None)
        threading.Thread(target=worker, daemon=True).start()
        win.after(100, poll); win.mainloop()
    except Exception as exc:
        LOG.exception('FAILED: %s', failure_details(exc))
        messagebox.showerror('Gagal', failure_details(exc), parent=win)
        return 1
    finally:
        win.destroy()
    return exit_code

def main(argv=None):
    configure_logging()
    LOG.info('START: bridge_bms terminal progress enabled')
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input', nargs='?', help='KMZ input; omitted = Tkinter workflow')
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--pkrms', nargs='?', const=DEFAULT_PKRMS_PATH, type=Path,
                        help='Import PKRMS Excel; omit the path after this flag to use the default workbook')
    parser.add_argument('--title', help='Document heading')
    parser.add_argument('--pu-logo', type=Path, default=PU_LOGO)
    parser.add_argument('--tanjabbar-logo', type=Path, default=TANJABBAR_LOGO)
    parser.add_argument('--no-logos', action='store_true', help='Explicitly generate without header logos')
    parser.add_argument('--kmz-only', action='store_true')
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args(argv)
    if not args.input:
        return run_gui(args)
    print(execute_job(args, prepare_job(args)))
    return 0

if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as exc:
        LOG.exception('FAILED: %s', failure_details(exc))
        sys.exit(1)
