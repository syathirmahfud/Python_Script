import os
import zipfile
import tempfile
import shutil
from openpyxl import load_workbook, Workbook
from tkinter import Tk, filedialog, messagebox

# Hide Tkinter root window
Tk().withdraw()

# === Select Folder with Excel Files ===
folder = filedialog.askdirectory(title="Select Folder Containing Excel Files")
if not folder:
    messagebox.showinfo("Cancelled", "No folder selected.")
    exit()

# === Select Output Save Location ===
output_file = filedialog.asksaveasfilename(
    title="Save Combined Excel As",
    defaultextension=".xlsx",
    filetypes=[("Excel Files", "*.xlsx")]
)
if not output_file:
    messagebox.showinfo("Cancelled", "No output file selected.")
    exit()

# === Helpers ===
def try_load_workbook(path):
    """Try to load workbook normally; return workbook or raise exception."""
    return load_workbook(path, data_only=True)

def repair_xlsx_remove_bad_media(path):
    """
    Try to create a repaired temp xlsx by skipping unreadable members (bad media).
    Returns path to repaired temp file (caller should delete it), or None on failure.
    """
    try:
        bad_names = []
        with zipfile.ZipFile(path, 'r') as zin:
            namelist = zin.namelist()

            # Try reading each member; mark unreadable entries
            readable = []
            for name in namelist:
                try:
                    _ = zin.read(name)  # attempt to read bytes
                    readable.append(name)
                except Exception:
                    bad_names.append(name)
                    # continue, we will skip this member when writing
            # If nothing is bad, no need to repair
            if not bad_names:
                return None

            # Build repaired zip without bad entries
            tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
            os.close(tmp_fd)
            with zipfile.ZipFile(tmp_path, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
                for name in readable:
                    data = zin.read(name)
                    zout.writestr(name, data)
            return tmp_path
    except Exception:
        return None

# === Create Output Workbook ===
wb_out = Workbook()
ws_out = wb_out.active
ws_out.title = "Merged"

start_row = 3
processed = []
skipped = []
repaired_files = []

# Loop all files
for filename in sorted(os.listdir(folder)):
    # skip temporary or non-xlsx/xlsm
    if filename.startswith("~$"):
        print("Skipping temp file:", filename)
        continue
    if not (filename.lower().endswith(".xlsx") or filename.lower().endswith(".xlsm")):
        continue

    filepath = os.path.join(folder, filename)
    print("Processing:", filename)
    wb = None
    repaired_temp = None
    try:
        # First attempt to open normally
        wb = try_load_workbook(filepath)
    except Exception as e:
        print(f"Normal load failed for {filename}: {e}")
        # Attempt to repair by removing bad media entries
        repaired_temp = repair_xlsx_remove_bad_media(filepath)
        if repaired_temp:
            try:
                print("Attempting to load repaired copy for", filename)
                wb = try_load_workbook(repaired_temp)
                repaired_files.append(filename)
            except Exception as e2:
                print("Repair load failed for", filename, "->", e2)
                wb = None
        else:
            print("Could not auto-repair", filename)

    if wb is None:
        print("Skipping file due to load error:", filename)
        skipped.append(filename)
        # cleanup repaired temp if created
        if repaired_temp and os.path.exists(repaired_temp):
            os.remove(repaired_temp)
        continue

    try:
        ws = wb.active
        # Detect last data row robustly by checking column B..L
        max_row = ws.max_row
        # move up if last rows are empty across B:L
        while max_row >= start_row:
            row_cells = [ws.cell(row=max_row, column=col).value for col in range(2, 13)]
            if any(c is not None for c in row_cells):
                break
            max_row -= 1

        if max_row < start_row:
            print("No data region found in", filename, "- skipping")
            skipped.append(filename)
        else:
            # Copy B3:L(max_row) and prepend filename at column A
            for row in ws.iter_rows(min_row=start_row, max_row=max_row, min_col=2, max_col=12):
                values = [cell.value for cell in row]
                out_row = [filename] + values
                ws_out.append(out_row)
            processed.append(filename)
            print("Appended rows from", filename)
    except Exception as e:
        print("Error while reading sheet from", filename, "->", e)
        skipped.append(filename)
    finally:
        # cleanup repaired temp
        if repaired_temp and os.path.exists(repaired_temp):
            os.remove(repaired_temp)

# Save output
wb_out.save(output_file)

# Show result summary
summary = f"Done.\nProcessed: {len(processed)} files\nSkipped: {len(skipped)} files\nRepaired: {len(repaired_files)} files\n\nSaved to:\n{output_file}"
print(summary)
if skipped:
    print("Skipped list:\n", "\n".join(skipped))
messagebox.showinfo("Merge result", summary)
